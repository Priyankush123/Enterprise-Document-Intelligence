"""
Excel Loader

Loads Excel workbooks (.xlsx) and converts them into
the standard Document schema.

Each worksheet is treated as one page.
"""

from pathlib import Path

from openpyxl import load_workbook

from src.ingestion.base_loader import BaseLoader
from src.ingestion.document_schema import (
    Document,
    DocumentMetadata,
    Page,
)


class ExcelLoader(BaseLoader):
    """
    Loader for Microsoft Excel files.
    """

    SUPPORTED_EXTENSIONS = [".xlsx"]

    def load(self, file_path: str) -> Document:

        excel_path = self.validate_file(file_path)

        workbook = load_workbook(
            excel_path,
            data_only=True
        )

        pages = []

        for index, sheet in enumerate(workbook.worksheets, start=1):

            lines = []

            for row in sheet.iter_rows(values_only=True):

                values = [
                    str(cell).strip()
                    for cell in row
                    if cell is not None
                ]

                if values:
                    lines.append(" | ".join(values))

            pages.append(
                Page(
                    page_number=index,
                    text="\n".join(lines)
                )
            )

        metadata = DocumentMetadata(
            title=excel_path.stem,
            page_count=len(pages),
        )

        return Document(
            filename=excel_path.name,
            filetype="xlsx",
            pages=pages,
            metadata=metadata,
        )