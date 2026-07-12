"""
DOCX Loader

Loads Microsoft Word (.docx) files and converts them into
the standard Document schema.
"""

from datetime import datetime
from pathlib import Path

from docx import Document as DocxDocument

from src.ingestion.base_loader import BaseLoader
from src.ingestion.document_schema import (
    Document,
    DocumentMetadata,
    Page,
)


class DOCXLoader(BaseLoader):
    """
    Loader for DOCX documents.
    """

    SUPPORTED_EXTENSIONS = [".docx"]

    def load(self, file_path: str) -> Document:
        """
        Load a DOCX document.
        """

        docx_path = self.validate_file(file_path)

        doc = DocxDocument(docx_path)

        paragraphs = []

        for para in doc.paragraphs:

            text = para.text.strip()

            if text:
                paragraphs.append(text)

        full_text = "\n\n".join(paragraphs)

        page = Page(
            page_number=1,
            text=full_text
        )

        metadata = DocumentMetadata(
            title=Path(file_path).stem,
            author=doc.core_properties.author,
            subject=doc.core_properties.subject,
            creator=doc.core_properties.author,
            created_at=doc.core_properties.created,
            modified_at=doc.core_properties.modified,
            page_count=1,
        )

        return Document(
            filename=docx_path.name,
            filetype="docx",
            pages=[page],
            metadata=metadata,
        )