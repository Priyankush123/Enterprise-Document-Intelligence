"""
PPTX Loader

Loads Microsoft PowerPoint presentations (.pptx)
and converts them into the standard Document schema.
"""

from pathlib import Path
from pptx import Presentation

from src.ingestion.base_loader import BaseLoader
from src.ingestion.document_schema import (
    Document,
    DocumentMetadata,
    Page,
)


class PPTXLoader(BaseLoader):
    """
    Loader for PowerPoint presentations.
    """

    SUPPORTED_EXTENSIONS = [".pptx"]

    def load(self, file_path: str) -> Document:

        ppt_path = self.validate_file(file_path)

        presentation = Presentation(ppt_path)

        pages = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):

            texts = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        texts.append(text)

            page = Page(
                page_number=slide_number,
                text="\n".join(texts),
            )

            pages.append(page)

        metadata = DocumentMetadata(
            title=Path(file_path).stem,
            author=presentation.core_properties.author,
            subject=presentation.core_properties.subject,
            creator=presentation.core_properties.author,
            created_at=presentation.core_properties.created,
            modified_at=presentation.core_properties.modified,
            page_count=len(pages),
        )

        return Document(
            filename=ppt_path.name,
            filetype="pptx",
            pages=pages,
            metadata=metadata,
        )